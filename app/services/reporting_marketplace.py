"""Marketplace reporting: Energy Schedules CSV + compliance CSV → hourly charts and LMP stats."""
from __future__ import annotations

import csv
import io
import math
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from statistics import mean
from typing import Any

from app.services.nomination_accuracy import (
    N_INTERVALS,
    fill_rtd_actual,
    load_mq_del_mw_from_xlsx,
    parse_compliance_csv,
)

# Compliance window aligned with nomination accuracy paste (MPI rows 05:05–19:00)
T_START = time(5, 5)
T_END = time(19, 0)

# Interval-end hours 6..18 = hourly LMP buckets from 6:00 through 18:00 (6AM–6PM display), used for average price.
_HOURS_LMP_AVG_6AM_6PM = range(6, 19)


def _parse_dt_interval_end(raw: str) -> datetime | None:
    raw = (raw or "").strip()
    if not raw:
        return None
    for fmt in ("%m/%d/%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def parse_market_result_energy_schedules(content: bytes) -> list[tuple[datetime, float, float]]:
    """
    MPI Market Result - Energy Schedules CSV.
    Returns rows ``(interval_end, mw, lmp)`` (MW = day-ahead / schedule MW, LMP).
    """
    text = content.decode("utf-8-sig", errors="replace")
    f = io.StringIO(text)
    reader = csv.reader(f)
    rows = list(reader)
    if not rows:
        return []
    header = [str(h).strip().lower() for h in rows[0]]

    def col(*names: str) -> int:
        for i, h in enumerate(header):
            for n in names:
                if n in h:
                    return i
        return -1

    i_end = col("interval end")
    i_mw = col("mw")
    i_lmp = col("lmp")
    if i_end < 0 or i_mw < 0 or i_lmp < 0:
        raise ValueError("Market Result CSV must have Interval End, MW, and LMP columns.")

    out: list[tuple[datetime, float, float]] = []
    for row in rows[1:]:
        if len(row) <= max(i_end, i_mw, i_lmp):
            continue
        dt = _parse_dt_interval_end(row[i_end])
        if dt is None:
            continue
        try:
            mw = float(str(row[i_mw]).replace(",", "").strip() or 0)
            lmp = float(str(row[i_lmp]).replace(",", "").strip() or 0)
        except ValueError:
            continue
        out.append((dt, mw, lmp))
    return out


def dominant_day_from_market_result_rows(rows: list[tuple[datetime, float, float]]) -> date | None:
    if not rows:
        return None
    counts: dict[date, int] = defaultdict(int)
    for dt, _, _ in rows:
        # Trading day for interval-ending-on-day rows
        d = dt.date()
        counts[d] += 1
    return max(counts, key=counts.get) if counts else None


def dominant_day_from_market_result_bytes(content: bytes) -> date | None:
    rows = parse_market_result_energy_schedules(content)
    return dominant_day_from_market_result_rows(rows)


def lmp_average_excel_rows_e7_through_e19(content: bytes) -> float:
    """
    Average LMP for the 13 data rows that fall on Excel rows 7–19 (1-based file lines:
    line 1 = header, so **file lines 7–19** = data row indices 5..17).
    Matches “LMP column E7:E19” when the export matches the standard MPI layout.
    """
    text = content.decode("utf-8-sig", errors="replace")
    f = io.StringIO(text)
    reader = csv.reader(f)
    rows = list(reader)
    if len(rows) < 19:
        raise ValueError("Market Result CSV has too few rows for LMP rows 7–19.")
    header = [str(h).strip().lower() for h in rows[0]]

    def col(*names: str) -> int:
        for i, h in enumerate(header):
            for n in names:
                if n in h:
                    return i
        return -1

    i_lmp = col("lmp")
    if i_lmp < 0:
        raise ValueError("Market Result CSV has no LMP column.")
    data = rows[1:]
    if len(data) < 18:
        raise ValueError("Not enough data rows for LMP rows 7–19.")
    slice_rows = data[5:18]
    lmps: list[float] = []
    for r in slice_rows:
        if len(r) <= i_lmp:
            continue
        try:
            lmps.append(float(str(r[i_lmp]).replace(",", "").strip() or 0))
        except ValueError:
            continue
    if not lmps:
        raise ValueError("Could not read LMP values for rows 7–19.")
    return float(mean(lmps))


def _compliance_rows_for_day(
    parsed: list[tuple[datetime, float, float]], day: date
) -> list[tuple[datetime, float, float]]:
    out: list[tuple[datetime, float, float]] = []
    for dt, rtd, act in parsed:
        if dt.date() != day:
            continue
        t = dt.time()
        if t < T_START or t > T_END:
            continue
        out.append((dt, rtd, act))
    out.sort(key=lambda x: x[0])
    return out


def _hourly_from_compliance(
    parsed: list[tuple[datetime, float, float]], day: date
) -> dict[int, tuple[float, float]]:
    """Hour -> (mean_rtd, mean_actual) from all 5-min samples in that clock hour on ``day``."""
    buckets: dict[int, list[tuple[float, float]]] = defaultdict(list)
    for dt, rtd, act in parsed:
        if dt.date() != day:
            continue
        buckets[dt.hour].append((rtd, act))
    out: dict[int, tuple[float, float]] = {}
    for h, pairs in buckets.items():
        if not pairs:
            continue
        out[h] = (
            float(mean([p[0] for p in pairs])),
            float(mean([p[1] for p in pairs])),
        )
    return out


def _hourly_he_from_compliance(
    parsed: list[tuple[datetime, float, float]], day: date, he_min: int, he_max: int
) -> dict[int, tuple[float, float]]:
    """
    Hour-ending (HE) buckets aligned with Excel / trading reports: **HE N** = intervals whose end
    time falls in ``((N-1):00, N:00]`` (e.g. HE 7 = 06:05 … 07:00), intersected with the compliance
    window. Labels like “6:00 AM” on the reference chart correspond to **HE 6**, not naive ``dt.hour``.
    """
    out: dict[int, tuple[float, float]] = {}
    for he in range(he_min, he_max + 1):
        lo = time(he - 1, 0, 0)
        hi = time(he, 0, 0)
        pairs: list[tuple[float, float]] = []
        for dt, rtd, act in parsed:
            if dt.date() != day:
                continue
            t = dt.time()
            if t < T_START or t > T_END:
                continue
            if not (lo < t <= hi):
                continue
            pairs.append((rtd, act))
        if pairs:
            out[he] = (
                float(mean([p[0] for p in pairs])),
                float(mean([p[1] for p in pairs])),
            )
    return out


def _market_maps_for_day(
    rows: list[tuple[datetime, float, float]], day: date
) -> tuple[dict[int, float], dict[int, float]]:
    """Hour -> MW and hour -> LMP for intervals ending on ``day`` (and 00:00 next day as hour 24)."""
    mw_h: dict[int, float] = {}
    lmp_h: dict[int, float] = {}
    next_mid = datetime.combine(day + timedelta(days=1), time(0, 0))
    for dt, mw, lmp in rows:
        if dt.date() == day:
            mw_h[dt.hour] = mw
            lmp_h[dt.hour] = lmp
        elif dt == next_mid:
            mw_h[24] = mw
            lmp_h[24] = lmp
    return mw_h, lmp_h


def _day_ahead_mw_linear(dt: datetime, mw_map: dict[int, float]) -> float | None:
    """
    Day-ahead schedule MW from Energy Schedules (one MW per clock-hour ``Interval End``).

    Excel line charts typically draw **straight segments** between points. MPI exports give one
    value per hour; repeating that constant for every 5-minute row produces a **step** (ladder).
    We **linearly interpolate** between consecutive hourly MW values at each interval-end time so
    the series matches that slanted segment look (same source data, different sampling).
    """
    if not mw_map:
        return None
    t = dt.hour + dt.minute / 60.0 + dt.second / 3600.0
    h0 = int(math.floor(t + 1e-9))
    frac = t - float(h0)
    if frac < 0:
        frac = 0.0
    if frac > 1:
        frac = 1.0
    v0 = mw_map.get(h0)
    v1 = mw_map.get(h0 + 1)
    if v0 is None and v1 is None:
        return None
    if v0 is None:
        return float(v1)
    if v1 is None:
        return float(v0)
    return float(v0) * (1.0 - frac) + float(v1) * frac


def build_marketplace_chart_payload(
    compliance_bytes: bytes,
    trade_day: date,
    market_bytes: bytes | None = None,
    mirf_mq_xlsx_bytes: bytes | None = None,
) -> dict[str, Any]:
    """
    Build JSON-serializable chart data from stored MPI compliance.

    If ``market_bytes`` is set, adds day-ahead MW and LMP from the **Market Result - Energy
    Schedules** CSV. **Average price (PHP)** is ``mean(LMP for interval-end hours 6..18) / 1000``,
    aligned with the hourly chart. The hourly **Actual MW** series uses **hour-ending (HE)**
    buckets ``((HE-1):00, HE:00]`` (same as Excel’s “hour ending” columns), not naive ``dt.hour``.
    Day-ahead MW on the dispatch chart is **linearly interpolated** between hourly schedule values.

    If ``market_bytes`` is omitted, returns the same shape with market fields ``null`` and
    ``partial: True`` (MPI-only charts). If ``mirf_mq_xlsx_bytes`` is set, day-ahead MW is
    filled from MIRF MQ (288×5-min) even while remaining partial (no LMP).
    """
    parsed = parse_compliance_csv(compliance_bytes)
    # Hourly combo chart: hour-ending HE 6..18 (6AM–6PM) to match Excel “hour ending” columns.
    hourly_he = _hourly_he_from_compliance(parsed, trade_day, 6, 18)

    mw_map: dict[int, float] = {}
    lmp_map: dict[int, float] = {}
    mq_5min: list[float] | None = None
    lmp_avg_e7_e19: float | None = None
    lmp_disp: float | None = None
    partial = market_bytes is None

    lmp_avg_hourly_6am_6pm: float | None = None
    average_price_php: float | None = None
    if market_bytes:
        mrows = parse_market_result_energy_schedules(market_bytes)
        mw_map, lmp_map = _market_maps_for_day(mrows, trade_day)
        lmp_avg_e7_e19 = lmp_average_excel_rows_e7_through_e19(market_bytes)
        lmp_disp = round(lmp_avg_e7_e19 / 1000.0, 3)
        hourly_lmps = [lmp_map[h] for h in _HOURS_LMP_AVG_6AM_6PM if h in lmp_map]
        if hourly_lmps:
            lmp_avg_hourly_6am_6pm = float(mean(hourly_lmps))
            average_price_php = round(lmp_avg_hourly_6am_6pm / 1000.0, 3)
    if mirf_mq_xlsx_bytes:
        # MIRF MQ is the plant's real submitted 5-minute schedule (DEL); prefer it as day-ahead
        # MW over the Market Result CSV's hourly figures linearly interpolated to 5-minute steps.
        mq_5min, _sheet = load_mq_del_mw_from_xlsx(mirf_mq_xlsx_bytes)

    win_rows = _compliance_rows_for_day(parsed, trade_day)
    dispatch: list[dict[str, Any]] = []

    def _idx_from_interval_end(dt: datetime) -> int | None:
        m = dt.hour * 60 + dt.minute + dt.second / 60.0
        if m % 5.0 > 1e-6:
            return None
        idx = int(round((m - 5) / 5))
        return idx if 0 <= idx < 288 else None

    def _mq_at(dt: datetime) -> float | None:
        if not mq_5min:
            return None
        idx = _idx_from_interval_end(dt)
        if idx is None:
            return None
        try:
            return float(mq_5min[idx])
        except (TypeError, ValueError):
            return None

    def _mq_hour_ending_avg(he: int) -> float | None:
        # HE N = interval-end times in ((N-1):00, N:00], which maps to 12×5-min slots.
        if not mq_5min:
            return None
        if he < 1 or he > 24:
            return None
        start_idx = (he - 1) * 12
        end_idx = start_idx + 12  # exclusive
        if start_idx < 0 or end_idx > 288:
            return None
        vals = mq_5min[start_idx:end_idx]
        return float(mean(vals)) if vals else None

    for i, (dt, rtd, act) in enumerate(win_rows, start=1):
        da = None
        if mq_5min:
            da = _mq_at(dt)
        elif mw_map:
            da = _day_ahead_mw_linear(dt, mw_map)
        dispatch.append(
            {
                "i": i,
                "interval_end": dt.isoformat(sep=" "),
                "rtd_mw": rtd,
                "actual_mw": act,
                "day_ahead_mw": da,
            }
        )

    hourly_chart: list[dict[str, Any]] = []
    actual_for_6am_6pm: list[float] = []
    for he in range(6, 19):
        rtd_m, act_m = hourly_he.get(he, (None, None))
        if act_m is not None:
            actual_for_6am_6pm.append(act_m)
        hourly_chart.append(
            {
                "hour_ending": he,
                "hour": he,
                "label": _hour_label_ampm(he),
                "rtd_mw_avg": rtd_m,
                "actual_mw_avg": act_m,
                "day_ahead_mw": _mq_hour_ending_avg(he) if mq_5min else (mw_map.get(he) if mw_map else None),
                "lmp": lmp_map.get(he) if lmp_map else None,
            }
        )

    actual_avg_6am_6pm = round(float(mean(actual_for_6am_6pm)), 4) if actual_for_6am_6pm else None
    actual_mwh_6am_6pm = round(float(sum(actual_for_6am_6pm)), 3) if actual_for_6am_6pm else None
    total_actual_mwh_hint = None
    if win_rows:
        total_actual_mwh_hint = float(sum(act for _, _, act in win_rows) / 12.0)

    out: dict[str, Any] = {
        "trade_day": trade_day.isoformat(),
        "partial": partial,
        "has_day_ahead": bool(mw_map) or bool(mq_5min),
        "day_ahead_source": "mirf_mq" if mq_5min else ("market_schedule" if mw_map else None),
        "lmp_average_e7_e19": lmp_avg_e7_e19,
        "lmp_average_e7_e19_display": lmp_disp,
        "lmp_average_hourly_6am_6pm": lmp_avg_hourly_6am_6pm,
        "average_price_php": average_price_php,
        "actual_dispatch_avg_mw_6am_6pm": actual_avg_6am_6pm,
        "actual_dispatch_mwh_6am_6pm": actual_mwh_6am_6pm,
        "actual_dispatch_mwh_compliance_window": total_actual_mwh_hint,
        "dispatch_series": dispatch,
        "hourly_6am_6pm": hourly_chart,
    }
    if partial and mq_5min:
        out["partial_message"] = (
            "MPI compliance is loaded from the database. Day-ahead MW is loaded from MIRF MQ "
            "(stored from Nomination Accuracy backfill). Upload Market Result - Energy Schedules "
            "for this trade day to add LMP and hourly average price."
        )
    return out


def _interval_index_for_time(t: time) -> int:
    m = t.hour * 60 + t.minute
    return int(round((m - 5) / 5))


def build_marketplace_xlsx_export(
    compliance_bytes: bytes,
    trade_day: date,
    market_bytes: bytes | None = None,
    mirf_mq_xlsx_bytes: bytes | None = None,
) -> bytes:
    """
    Render one trade day's MPI compliance (+ optional Market Result / MIRF MQ day-ahead) as an
    ``.xlsx`` workbook matching the plant's own ``RTD -Actual -Day Ahead`` template: a
    5-minute INTERVAL / RTD (MW) / ACTUAL DISPATCH (MW) / Day Ahead grid restricted to the
    compliance window (05:05-19:00) plus an embedded "RTD VS ACTUAL" line chart, so a day's
    numbers can be archived/shared in the same layout as the manually-built workbook instead
    of a CSV.
    """
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    parsed = parse_compliance_csv(compliance_bytes)
    rtd, actual, _n_applied = fill_rtd_actual(parsed, trade_day, T_START, T_END)

    mw_map: dict[int, float] = {}
    lmp_map: dict[int, float] = {}
    mq_5min: list[float] | None = None
    if market_bytes:
        mrows = parse_market_result_energy_schedules(market_bytes)
        mw_map, lmp_map = _market_maps_for_day(mrows, trade_day)
    if mirf_mq_xlsx_bytes:
        mq_5min, _sheet = load_mq_del_mw_from_xlsx(mirf_mq_xlsx_bytes)

    def _day_ahead_at(i: int) -> float | None:
        # MIRF MQ is the plant's real submitted 5-minute schedule; prefer it over the Market
        # Result CSV's hourly figures linearly interpolated to 5-minute steps (an estimate).
        if mq_5min:
            return float(mq_5min[i])
        minute_of_day = 5 * (i + 1)
        hh = (minute_of_day // 60) % 24
        mm = minute_of_day % 60
        base_day = trade_day + timedelta(days=1) if minute_of_day >= 1440 else trade_day
        dt = datetime.combine(base_day, time(hh, mm))
        if mw_map:
            return _day_ahead_mw_linear(dt, mw_map)
        return None

    win_start_idx = _interval_index_for_time(T_START)
    win_end_idx = _interval_index_for_time(T_END)
    n_window = win_end_idx - win_start_idx + 1
    win_start_row = 3
    win_end_row = 2 + n_window

    wb = Workbook()
    ws = wb.active
    ws.title = "RTD -Actual -Day Ahead"
    ws.sheet_view.showGridLines = False

    header_font = Font(bold=True, size=12)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(size=12)
    data_align = Alignment(horizontal="center", vertical="center")
    table_fill = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
    thin = Side(style="thin", color="FFB0AFAF")
    box_border = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws.merge_cells("C1:C2")
    ws.merge_cells("D1:D2")
    ws.merge_cells("E1:E2")
    ws.merge_cells("F1:F2")
    ws["C1"] = "INTERVAL"
    ws["D1"] = "RTD\n(MW)"
    ws["E1"] = "ACTUAL DISPATCH (MW)"
    ws["F1"] = "Day Ahead"
    ws["B2"] = "Time"
    ws["H2"] = "INTERVAL"
    ws["I2"] = "Average Price Hourly"
    for coord in ("B2", "C1", "D1", "E1", "F1", "H2", "I2"):
        ws[coord].font = header_font
        ws[coord].alignment = header_align
        ws[coord].fill = table_fill
        ws[coord].border = box_border
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 14.5
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 17.8
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 14
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 30

    # B/D/E/F: 5-minute rows for the compliance window ONLY (05:05-19:00), filled and
    # bordered throughout so the sheet reads as one solid table.
    for i in range(win_start_idx, win_end_idx + 1):
        row = 3 + (i - win_start_idx)
        minute_of_day = 5 * (i + 1)
        hh = (minute_of_day // 60) % 24
        mm = minute_of_day % 60
        for col in (2, 3, 4, 5, 6):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.alignment = data_align
            cell.fill = table_fill
            cell.border = box_border
        ws.cell(row=row, column=2, value=(minute_of_day - 1) // 60)
        ws.cell(row=row, column=3, value=time(hh, mm)).number_format = "h:mm"
        ws.cell(row=row, column=4, value=round(rtd[i], 3))
        ws.cell(row=row, column=5, value=round(actual[i], 3))
        da = _day_ahead_at(i)
        if da is not None:
            ws.cell(row=row, column=6, value=round(da, 3))

    # H/I: hourly LMP reference table (hour-ending 6..18) + its average.
    if lmp_map:
        he_hours = list(range(6, 19))
        for k, he in enumerate(he_hours):
            r = win_start_row + k
            ws.cell(row=r, column=8, value=time(he % 24, 0)).number_format = "h:mm"
            if he in lmp_map:
                ws.cell(row=r, column=9, value=round(lmp_map[he], 2))
            for col in (8, 9):
                cell = ws.cell(row=r, column=col)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = box_border
        avg_row = win_start_row + len(he_hours)
        avg_cell = ws.cell(
            row=avg_row,
            column=9,
            value=f"=AVERAGE(I{win_start_row}:I{win_start_row + len(he_hours) - 1})",
        )
        avg_cell.font = Font(bold=True, size=12)
        avg_cell.alignment = data_align
        avg_cell.border = box_border

    ws.freeze_panes = "C3"

    chart = LineChart()
    chart.title = "RTD VS ACTUAL"
    chart.width = 22
    chart.height = 10
    chart.style = 2
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.tickLblSkip = 12  # ~hourly labels instead of every 5-min point, to stop overlap
    rtd_ref = Reference(ws, min_col=4, min_row=win_start_row, max_row=win_end_row)
    act_ref = Reference(ws, min_col=5, min_row=win_start_row, max_row=win_end_row)
    cats = Reference(ws, min_col=3, min_row=win_start_row, max_row=win_end_row)
    chart.add_data(rtd_ref, titles_from_data=False)
    chart.add_data(act_ref, titles_from_data=False)
    chart.series[0].tx = SeriesLabel(v="RTD (Market DOT)")
    chart.series[1].tx = SeriesLabel(v="Actual Dispatch")
    for s in chart.series:
        s.marker.symbol = "none"
        s.smooth = False
    chart.set_categories(cats)
    # Anchor right next to the main table (where the hourly rollup block used to sit),
    # top-aligned with the header instead of far below the whole table.
    ws.add_chart(chart, "K1")

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _hour_label_ampm(h: int) -> str:
    if h == 0:
        return "12AM"
    if h < 12:
        return str(h) + "AM"
    if h == 12:
        return "12PM"
    return str(h - 12) + "PM"

